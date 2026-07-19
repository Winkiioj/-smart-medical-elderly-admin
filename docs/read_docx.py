import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document

d = Document(r'D:\ProgramData\work\-smart-medical-elderly-admin\docs\SMECMS_测试用例.docx')
print("=== 段落 ===")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        print(f"P{i}: {t[:120]}")

print(f"\n=== 表格({len(d.tables)}个) ===")
for i, t in enumerate(d.tables):
    r0c0 = t.cell(0, 0).text[:50] if t.rows and t.columns else "?"
    row_count = len(t.rows)
    print(f"Table{i}: {row_count}行, 首格={r0c0}")
    if i == 0 or i == 1 or i == 2:
        for j, r in enumerate(t.rows):
            print(f"  行{j}: {[c.text[:60] for c in r.cells]}")
